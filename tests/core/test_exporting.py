"""تست خروجی CSV و Excel، بدون دیتابیس و مدل."""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook

from rasam import cli
from rasam.core import exporting
from rasam.core.graph import Answer


ROWS = [{"device": "in1", "avg": 28.6}, {"device": "in2", "max": 31.2}]


def test_csv_contains_all_columns_and_rows():
    assert exporting.to_csv(ROWS).splitlines() == [
        "device,avg,max", "in1,28.6,", "in2,,31.2"]


def test_excel_contains_all_columns_and_rows():
    workbook = load_workbook(io.BytesIO(exporting.to_excel(ROWS)))
    assert list(workbook.active.values) == [
        ("device", "avg", "max"), ("in1", 28.6, None), ("in2", None, 31.2)]


def test_cli_writes_requested_csv(monkeypatch, tmp_path):
    path = tmp_path / "result.csv"
    monkeypatch.setattr(cli, "answer", lambda _: Answer(question="q", rows=ROWS))
    assert cli.main(["--csv", str(path), "q"]) == 0
    assert path.read_text(encoding="utf-8") == exporting.to_csv(ROWS)

def test_a_cell_that_excel_would_run_as_a_formula_is_neutralized():
    """محتوای دیتابیس غیرقابل‌اعتماد است.

    سلولی که با `=` شروع شود، Excel موقع باز کردن فایل *اجرایش* می‌کند — یعنی
    متنی که کسی در یک فیلد دیتابیس نوشته، روی کامپیوتر گیرنده‌ی فایل فرمان
    می‌شود.
    """
    rows = [{"note": "=HYPERLINK(\"http://evil\",\"click\")", "other": "@SUM(A1)"}]
    text = exporting.to_csv(rows)
    assert "'=HYPERLINK" in text
    assert "'@SUM" in text

    sheet = load_workbook(io.BytesIO(exporting.to_excel(rows))).active
    assert sheet.cell(row=2, column=1).value.startswith("'=")


def test_a_negative_number_is_still_a_number_not_text():
    """`-27.31` عدد منفی است نه فرمول؛ متن‌کردنش یعنی خراب‌کردن جواب."""
    rows = [{"avg": -27.31}]
    assert exporting.to_csv(rows) == "avg\n-27.31\n"
    sheet = load_workbook(io.BytesIO(exporting.to_excel(rows))).active
    assert sheet.cell(row=2, column=1).value == -27.31
